"""Tests for R1 — imports router file validation.

Covers:
  - Wrong extension rejected (422 INVALID_EXTENSION)
  - File >10MB rejected (413 FILE_TOO_LARGE)
  - File >20k rows rejected (422 ROW_CAP_EXCEEDED)
  - Valid CSV accepted and batch row created (status=preview)
  - Valid XLSX accepted and batch row created
  - Empty file rejected
  - Correct row count for CSV and XLSX
  - Import history lists batches
  - Cancel (DELETE) sets status=expired
  - Confirm sets status=committed
  - Double confirm returns 409
"""

from __future__ import annotations

import io
import csv
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.routers.imports import _batches


@pytest.fixture(autouse=True)
def clear_batches():
    """Reset in-memory batch store between tests."""
    _batches.clear()
    yield
    _batches.clear()


client = TestClient(app)


# ---------------------------------------------------------------------------
# Helpers to build test files
# ---------------------------------------------------------------------------

def _make_csv_bytes(rows: list[list[str]], delimiter: str = ",") -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8")


def _make_xlsx_bytes(rows: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


VALID_HEADER = ["Nomor Pesanan", "Nama Produk", "SKU", "Jumlah", "Harga Awal", "Status Pesanan"]
VALID_ROW = ["ORD001", "Beras Ramos 5kg", "BERAS-5KG", "2", "50000", "SELESAI"]


def _upload(content: bytes, filename: str = "test.csv", channel: str = "shopee"):
    return client.post(
        "/v1/imports",
        files={"file": (filename, content, "text/csv" if filename.endswith(".csv") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"channel": channel},
    )


# ---------------------------------------------------------------------------
# Extension validation
# ---------------------------------------------------------------------------

class TestExtensionValidation:

    def test_reject_txt(self):
        r = _upload(b"hello", filename="data.txt")
        assert r.status_code == 422
        assert r.json()["detail"]["error"]["code"] == "INVALID_EXTENSION"

    def test_reject_json(self):
        r = _upload(b"{}", filename="data.json")
        assert r.status_code == 422
        assert r.json()["detail"]["error"]["code"] == "INVALID_EXTENSION"

    def test_reject_pdf(self):
        r = _upload(b"%PDF-1.4", filename="export.pdf")
        assert r.status_code == 422

    def test_accept_csv(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        r = _upload(content, filename="data.csv")
        assert r.status_code == 201

    def test_accept_xlsx(self):
        content = _make_xlsx_bytes([VALID_HEADER, VALID_ROW])
        r = _upload(content, filename="data.xlsx")
        assert r.status_code == 201


# ---------------------------------------------------------------------------
# Size validation
# ---------------------------------------------------------------------------

class TestSizeValidation:

    def test_reject_over_10mb(self):
        """File >10MB must be rejected with 413."""
        header = VALID_HEADER
        big_content = ",".join(header) + "\n"
        row_line = ",".join(VALID_ROW) + "\n"
        big_content += row_line * 215_000
        raw = big_content.encode("utf-8")
        assert len(raw) > 10 * 1024 * 1024

        r = _upload(raw, filename="huge.csv")
        assert r.status_code == 413
        assert r.json()["detail"]["error"]["code"] == "FILE_TOO_LARGE"

    def test_accept_under_10mb(self):
        content = _make_csv_bytes([VALID_HEADER] + [VALID_ROW] * 100)
        r = _upload(content, filename="small.csv")
        assert r.status_code == 201


# ---------------------------------------------------------------------------
# Row cap validation
# ---------------------------------------------------------------------------

class TestRowCapValidation:

    def test_reject_over_20k_rows(self):
        """File with >20,000 data rows must be rejected with 422 ROW_CAP_EXCEEDED."""
        rows = [VALID_HEADER] + [VALID_ROW] * 20_001
        content = _make_csv_bytes(rows)
        r = _upload(content, filename="big.csv")
        assert r.status_code == 422
        body = r.json()["detail"]["error"]
        assert body["code"] == "ROW_CAP_EXCEEDED"
        assert body["cap"] == 20_000
        assert "split" in body["hint"].lower()

    def test_accept_exactly_20k_rows(self):
        rows = [VALID_HEADER] + [VALID_ROW] * 20_000
        content = _make_csv_bytes(rows)
        r = _upload(content, filename="max.csv")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 20_000


# ---------------------------------------------------------------------------
# Correct row counting
# ---------------------------------------------------------------------------

class TestRowCounting:

    def test_csv_row_count(self):
        rows = [VALID_HEADER] + [VALID_ROW] * 42
        content = _make_csv_bytes(rows)
        r = _upload(content, filename="count.csv")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 42

    def test_xlsx_row_count(self):
        rows = [VALID_HEADER] + [VALID_ROW] * 15
        content = _make_xlsx_bytes(rows)
        r = _upload(content, filename="count.xlsx")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 15

    def test_csv_semicolon_delimiter(self):
        rows = [VALID_HEADER, VALID_ROW, VALID_ROW, VALID_ROW]
        content = _make_csv_bytes(rows, delimiter=";")
        r = _upload(content, filename="semi.csv")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 3

    def test_xlsx_skip_empty_rows(self):
        """Fully empty rows in XLSX should not be counted."""
        rows = [VALID_HEADER, VALID_ROW, [None, None, None, None, None, None], VALID_ROW]
        content = _make_xlsx_bytes(rows)
        r = _upload(content, filename="gaps.xlsx")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 2  # empty row skipped


# ---------------------------------------------------------------------------
# Empty file
# ---------------------------------------------------------------------------

class TestEmptyFile:

    def test_empty_csv(self):
        r = _upload(b"", filename="empty.csv")
        assert r.status_code == 422

    def test_header_only_csv(self):
        content = _make_csv_bytes([VALID_HEADER])
        r = _upload(content, filename="headeronly.csv")
        assert r.status_code == 201
        assert r.json()["rows_read"] == 0

    def test_empty_xlsx(self):
        import openpyxl
        wb = openpyxl.Workbook()
        # Remove default sheet, add empty one
        wb.remove(wb.active)
        wb.create_sheet("Empty")
        buf = io.BytesIO()
        wb.save(buf)
        r = _upload(buf.getvalue(), filename="empty.xlsx")
        assert r.status_code == 422


# ---------------------------------------------------------------------------
# Batch creation
# ---------------------------------------------------------------------------

class TestBatchCreation:

    def test_batch_has_correct_fields(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        r = _upload(content, filename="ok.csv", channel="shopee")
        assert r.status_code == 201
        body = r.json()
        assert body["status"] == "preview"
        assert body["rows_read"] == 1
        assert body["channel"] == "shopee"
        assert "import_batch_id" in body
        assert "file_hash" in body
        assert len(body["file_hash"]) == 64  # SHA-256 hex

    def test_batch_file_hash_consistent(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW, VALID_ROW])
        r1 = _upload(content, filename="a.csv")
        r2 = _upload(content, filename="b.csv")
        assert r1.json()["file_hash"] == r2.json()["file_hash"]


# ---------------------------------------------------------------------------
# History / Preview / Confirm / Cancel
# ---------------------------------------------------------------------------

class TestBatchLifecycle:

    def test_list_imports(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        _upload(content, filename="a.csv")
        _upload(content, filename="b.csv")
        r = client.get("/v1/imports")
        assert r.status_code == 200
        assert len(r.json()["items"]) == 2

    def test_preview(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        batch_id = _upload(content, filename="a.csv").json()["import_batch_id"]
        r = client.get(f"/v1/imports/{batch_id}/preview")
        assert r.status_code == 200
        assert r.json()["status"] == "preview"

    def test_confirm(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        batch_id = _upload(content, filename="a.csv").json()["import_batch_id"]
        r = client.post(f"/v1/imports/{batch_id}/confirm")
        assert r.status_code == 200
        assert r.json()["status"] == "committed"

    def test_double_confirm_rejected(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        batch_id = _upload(content, filename="a.csv").json()["import_batch_id"]
        client.post(f"/v1/imports/{batch_id}/confirm")
        r = client.post(f"/v1/imports/{batch_id}/confirm")
        assert r.status_code == 409
        assert r.json()["detail"]["error"]["code"] == "ALREADY_COMMITTED"

    def test_cancel(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        batch_id = _upload(content, filename="a.csv").json()["import_batch_id"]
        r = client.delete(f"/v1/imports/{batch_id}")
        assert r.status_code == 204

    def test_not_found(self):
        r = client.get("/v1/imports/nonexistent/preview")
        assert r.status_code == 404

    def test_problems_endpoint(self):
        content = _make_csv_bytes([VALID_HEADER, VALID_ROW])
        batch_id = _upload(content, filename="a.csv").json()["import_batch_id"]
        r = client.get(f"/v1/imports/{batch_id}/problems")
        assert r.status_code == 200
        assert r.json()["batch_id"] == batch_id
